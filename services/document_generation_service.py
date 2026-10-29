# ============================================================
# شرح الملف: services/document_generation_service.py
# مركز موحّد لإنشاء الملفات من بيانات BuildOffice Pro.
# لا تحتوي الصفحات التشغيلية على منطق حفظ ملفات بعد الآن.
# ============================================================
"""Central file-generation services used exclusively by DocumentGenerationPage."""
from __future__ import annotations

import csv
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from core.config import generated_dir
from repositories.db import get_connection

OUTPUT_DIR = generated_dir() / "document_center"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def _timestamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")


def _write_text(filename: str, text: str) -> Path:
    path = OUTPUT_DIR / filename
    path.write_text(text, encoding="utf-8")
    return path


def _write_csv(filename: str, headers: list[str], rows: list[list[Any]]) -> Path:
    path = OUTPUT_DIR / filename
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)
    return path


def project_financial_report(project_id: str | None = None) -> tuple[str, list[list[Any]]]:
    conn = get_connection()
    try:
        sql = """SELECT p.project_id,p.project_name,p.budget_ht,p.budget_ttc,p.status,
            (SELECT COALESCE(SUM(amount),0) FROM expenses e WHERE e.project_id=p.project_id),
            (SELECT COALESCE(SUM(amount),0) FROM project_payments pp WHERE pp.project_id=p.project_id)
            FROM projects p"""
        args: tuple[Any, ...] = ()
        if project_id:
            sql += " WHERE p.project_id=?"
            args = (project_id,)
        rows = conn.execute(sql + " ORDER BY project_name", args).fetchall()
    finally:
        conn.close()
    lines = ["BUILDOffice Pro - PROJECT FINANCIAL REPORT", f"Generated: {datetime.now(timezone.utc):%Y-%m-%d %H:%M}", ""]
    data = []
    for r in rows:
        lines.append(f"{r['project_id']} | {r['project_name']} | HT {float(r['budget_ht'] or 0):,.2f} | Expenses {float(r[5] or 0):,.2f} | Received {float(r[6] or 0):,.2f} | {r['status']}")
        data.append([r['project_id'], r['project_name'], r['budget_ht'], r['budget_ttc'], r['status'], r[5], r[6]])
    return "\n".join(lines), data


def expense_report() -> tuple[str, list[list[Any]]]:
    conn = get_connection()
    try:
        rows = conn.execute("""SELECT e.expense_id,COALESCE(p.project_name,'General'),e.date,e.description,e.category,e.amount,
            COALESCE((SELECT SUM(amount) FROM expense_payments ep WHERE ep.expense_id=e.expense_id),0)
            FROM expenses e LEFT JOIN projects p ON p.project_id=e.project_id ORDER BY e.date DESC""").fetchall()
    finally:
        conn.close()
    lines = ["BUILDOffice Pro - EXPENSE REPORT", "ID | Project | Date | Description | Category | Amount | Paid"]
    data = []
    for r in rows:
        lines.append(f"{r[0]} | {r[1]} | {r[2]} | {r[3]} | {r[4] or '-'} | {float(r[5] or 0):,.2f} | {float(r[6] or 0):,.2f}")
        data.append(list(r))
    return "\n".join(lines), data


def attendance_report() -> tuple[str, list[list[Any]]]:
    conn = get_connection()
    try:
        rows = conn.execute("""SELECT a.date,s.full_name,COALESCE(p.project_name,'Administration'),a.status,a.working_hours,a.overtime_hours
            FROM attendance a JOIN staff s ON s.staff_id=a.staff_id LEFT JOIN projects p ON p.project_id=a.project_id ORDER BY a.date DESC""").fetchall()
    finally:
        conn.close()
    lines = ["BUILDOffice Pro - ATTENDANCE REPORT", "Date | Employee | Project | Status | Regular h | Overtime h"]
    data = []
    for r in rows:
        lines.append(f"{r[0]} | {r[1]} | {r[2]} | {r[3]} | {float(r[4] or 0):.2f} | {float(r[5] or 0):.2f}")
        data.append(list(r))
    return "\n".join(lines), data


def boq_report(project_id: str | None = None) -> tuple[str, list[list[Any]]]:
    conn = get_connection()
    try:
        sql = """SELECT p.project_name,m.item_number,m.description,m.unit,m.quantity,m.unit_price,m.amount
            FROM metre_items m JOIN projects p ON p.project_id=m.project_id"""
        args: tuple[Any, ...] = ()
        if project_id:
            sql += " WHERE p.project_id=?"
            args = (project_id,)
        rows = conn.execute(sql + " ORDER BY p.project_name,m.item_number", args).fetchall()
    finally:
        conn.close()
    lines = ["BUILDOffice Pro - BOQ / METRE REPORT", "Project | No. | Description | Unit | Qty | Unit price | Amount"]
    data = []
    for r in rows:
        lines.append(f"{r[0]} | {r[1] or '-'} | {r[2]} | {r[3] or '-'} | {float(r[4] or 0):.3f} | {float(r[5] or 0):,.2f} | {float(r[6] or 0):,.2f}")
        data.append(list(r))
    return "\n".join(lines), data


def generate_text_or_csv(kind: str, project_id: str | None = None, fmt: str = "TXT") -> list[Path]:
    builders = {
        "project_financial": project_financial_report,
        "expenses": expense_report,
        "attendance": attendance_report,
        "boq": boq_report,
    }
    if kind not in builders:
        raise ValueError(f"Unsupported document generator: {kind}")
    text, rows = builders[kind](project_id) if kind in ("project_financial", "boq") else builders[kind]()
    stamp = _timestamp()
    if fmt.upper() == "CSV":
        headers = {
            "project_financial": ["Project ID","Project","Budget HT","Budget TTC","Status","Expenses","Received"],
            "expenses": ["ID","Project","Date","Description","Category","Amount","Paid"],
            "attendance": ["Date","Employee","Project","Status","Regular Hours","Overtime Hours"],
            "boq": ["Project","No.","Description","Unit","Qty","Unit Price","Amount"],
        }[kind]
        return [_write_csv(f"{kind}_{stamp}.csv", headers, rows)]
    return [_write_text(f"{kind}_{stamp}.txt", text)]


def generate_project_calendar(project_id: str) -> Path:
    """Generate a simple Excel project-period calendar from the project dates."""
    conn = get_connection()
    try:
        row = conn.execute("SELECT project_id,project_name,start_date,end_date,duration_days,status FROM projects WHERE project_id=?", (project_id,)).fetchone()
    finally:
        conn.close()
    if not row:
        raise ValueError("Project not found.")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as exc:
        raise RuntimeError("Project calendar generation requires openpyxl.") from exc
    wb = Workbook()
    ws = wb.active
    if ws is None:
        raise RuntimeError("Failed to create the project calendar worksheet.")
    ws.title = "Project Calendar"
    ws.append(["BuildOffice Pro - Project Period Calendar"])
    ws.append(["Project ID", row["project_id"]]); ws.append(["Project", row["project_name"]]); ws.append(["Start", row["start_date"]]); ws.append(["End", row["end_date"]]); ws.append(["Duration (days)", row["duration_days"]]); ws.append(["Status", row["status"]])
    for cell in ws[1]: cell.font = Font(bold=True, size=14)
    for col in ("A","B"): ws.column_dimensions[col].width = 24 if col == "A" else 55
    for row_cells in ws.iter_rows():
        for cell in row_cells: cell.alignment = Alignment(vertical="top")
    path = OUTPUT_DIR / f"project_calendar_{row['project_id']}_{_timestamp()}.xlsx"; wb.save(path); return path
