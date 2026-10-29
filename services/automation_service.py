"""BuildOffice Pro automation kernel.

Provides shared smart defaults, computed fields, event-driven post-actions,
batch import parsing, and one-click report generation without changing the
existing database schema or page contracts.
"""
from __future__ import annotations

import csv
import json
import sqlite3
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from core.logging_config import logger
from database.database import get_connection

try:
    from openpyxl import load_workbook
except ImportError:  # optional dependency
    load_workbook = None


@dataclass(slots=True)
class ProjectDefaults:
    project_id: str | None = None
    project_name: str = ""
    authority: str = ""
    project_type: str = ""
    location: str = ""
    contract_date: str = ""
    start_date: str = ""
    duration_days: int = 0
    tva_percent: float = 19.0
    status: str = "Active"


@dataclass(slots=True)
class ExpenseDefaults:
    project_id: str | None = None
    date: str = ""
    category: str = ""
    payment_method: str = ""
    unit: str = "Lump Sum"
    quantity: float = 1.0


def today_iso() -> str:
    return datetime.now(timezone.utc).date().isoformat()


def _row_value(row: sqlite3.Row | None, key: str, default: Any = None) -> Any:
    if row is None:
        return default
    try:
        value = row[key]
    except (KeyError, IndexError):
        return default
    return default if value is None else value


def last_nonempty(conn: sqlite3.Connection, table: str, column: str, where: str = "", params: tuple[Any, ...] = ()) -> str:
    """Return the most recently used non-empty value for a trusted table/column."""
    sql = f"SELECT {column} FROM {table} WHERE {column} IS NOT NULL AND TRIM(CAST({column} AS TEXT)) <> ''"
    if where:
        sql += f" AND ({where})"
    sql += " ORDER BY rowid DESC LIMIT 1"
    try:
        row = conn.execute(sql, params).fetchone()
        return str(row[0]) if row and row[0] is not None else ""
    except sqlite3.Error:
        return ""


def project_defaults(conn: sqlite3.Connection | None = None, project_id: str | None = None, preferred_project_id: str | None = None) -> ProjectDefaults:
    """Build a project form from current date, company context, and history."""
    own = conn is None
    target_conn = conn or get_connection()
    try:
        target = project_id or preferred_project_id
        row = None
        if target:
            row = target_conn.execute("SELECT * FROM projects WHERE project_id=?", (target,)).fetchone()
        company = target_conn.execute('SELECT * FROM "company info" ORDER BY id LIMIT 1').fetchone()
        today = today_iso()
        if row:
            return ProjectDefaults(
                project_id=str(_row_value(row, "project_id", target)),
                project_name=str(_row_value(row, "project_name", "") or ""),
                authority=str(_row_value(row, "contracting_authority", "") or ""),
                project_type=str(_row_value(row, "project_type", "") or ""),
                location=str(_row_value(row, "location", "") or ""),
                contract_date=str(_row_value(row, "contract_date", today) or today),
                start_date=str(_row_value(row, "start_date", today) or today),
                duration_days=int(_row_value(row, "period", 0) or 0),
                tva_percent=float(_row_value(row, "tva_percent", 19) or 19),
                status=str(_row_value(row, "status", "Active") or "Active"),
            )

        recent = target_conn.execute(
            "SELECT contracting_authority, project_type, location, period, tva_percent FROM projects "
            "ORDER BY COALESCE(updated_at, created_at) DESC LIMIT 1"
        ).fetchone()
        return ProjectDefaults(
            authority=str(_row_value(recent, "contracting_authority", "") or "") or str(_row_value(company, "company name", "") or ""),
            project_type=str(_row_value(recent, "project_type", "") or ""),
            location=str(_row_value(recent, "location", "") or ""),
            contract_date=today,
            start_date=today,
            duration_days=int(_row_value(recent, "period", 0) or 0),
            tva_percent=float(_row_value(recent, "tva_percent", 19) or 19),
            status="Active",
        )
    finally:
        if own:
            target_conn.close()


def expense_defaults(conn: sqlite3.Connection, preferred_project_id: str | None = None) -> ExpenseDefaults:
    project_id = preferred_project_id
    if not project_id:
        row = conn.execute(
            "SELECT project_id FROM expenses WHERE project_id IS NOT NULL AND TRIM(project_id)<>'' "
            "ORDER BY COALESCE(updated_at, created_at) DESC LIMIT 1"
        ).fetchone()
        project_id = str(row[0]) if row else None
    return ExpenseDefaults(
        project_id=project_id,
        date=today_iso(),
        category=last_nonempty(conn, "expenses", "category"),
        payment_method=last_nonempty(conn, "expenses", "payment_method"),
        unit=last_nonempty(conn, "expenses", "unit") or "Lump Sum",
        quantity=1.0,
    )


def compute_expense_total(quantity: float, unit_price: float, mode: str, lump_sum: float = 0.0) -> float:
    if str(mode).strip().lower() in {"unit-based", "unit based", "unit"}:
        return round(float(quantity or 0) * float(unit_price or 0), 2)
    return round(float(lump_sum or 0), 2)


def compute_project_end(start_date: str, duration_days: int) -> str:
    try:
        start = datetime.strptime(str(start_date), "%Y-%m-%d").replace(tzinfo=timezone.utc).date()
    except ValueError:
        start = datetime.now(timezone.utc).date()
    return (start + timedelta(days=max(0, int(duration_days or 0)))).isoformat()


def payment_status(total: float, paid: float, due_date: str | None = None) -> str:
    total = max(0.0, float(total or 0))
    paid = max(0.0, float(paid or 0))
    if 0 < total <= paid:
        return "Paid"
    if paid > 0:
        return "Partial"
    if due_date and due_date < today_iso():
        return "Overdue"
    return "Unpaid"


def calculate_progress(executed: float, contract_value: float) -> float:
    return round((float(executed or 0) / float(contract_value)) * 100.0, 2) if contract_value else 0.0


def parse_rows(path: str | Path) -> tuple[list[str], list[list[Any]]]:
    """Parse CSV/XLSX once into a normalized header+rows representation."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".csv":
        with p.open("r", encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.reader(fh))
        return (rows[0] if rows else []), rows[1:]
    if suffix in {".xlsx", ".xlsm"}:
        if load_workbook is None:
            raise RuntimeError("Excel import requires openpyxl.")
        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb.active
        assert ws is not None
        values = [list(row) for row in ws.iter_rows(values_only=True)]
        wb.close()
        return (["" if x is None else str(x).strip() for x in values[0]] if values else []), values[1:]
    raise ValueError("Supported batch imports: CSV and XLSX.")


def normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    return "".join(ch for ch in text if ch.isalnum())


def suggest_mapping(headers: Iterable[Any], target_columns: Iterable[str]) -> dict[str, str]:
    src = {normalize_header(h): str(h) for h in headers}
    mapping: dict[str, str] = {}
    for target in target_columns:
        nt = normalize_header(target)
        if nt in src:
            mapping[target] = src[nt]
            continue
        for key, original in src.items():
            if nt in key or key in nt:
                mapping[target] = original
                break
    return mapping


def import_rows(table: str, headers: list[Any], rows: list[list[Any]], mapping: dict[str, str], *, dry_run: bool = False, conn: sqlite3.Connection | None = None) -> dict[str, Any]:
    """Safe batch insert for existing tables. It never creates or renames schema."""
    own = conn is None
    target = conn or get_connection()
    try:
        columns = [str(r[1]) for r in target.execute(f"PRAGMA table_info({table})").fetchall()]
        normalized_lookup = {normalize_header(c): c for c in columns}
        resolved: list[tuple[str, int]] = []
        header_lookup = {str(h): i for i, h in enumerate(headers)}
        for target_col, source_header in mapping.items():
            actual = normalized_lookup.get(normalize_header(target_col))
            if actual is not None and source_header in header_lookup:
                resolved.append((actual, header_lookup[source_header]))
        if not resolved:
            raise ValueError("No valid import columns were mapped.")

        insert_cols = [c for c, _ in resolved]
        placeholders = ",".join("?" for _ in insert_cols)
        sql = f"INSERT INTO {table} ({','.join(insert_cols)}) VALUES ({placeholders})"
        good, skipped = 0, []
        for line_no, row in enumerate(rows, start=2):
            values = [row[idx] if idx < len(row) else None for _, idx in resolved]
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            if dry_run:
                good += 1
                continue
            try:
                target.execute(sql, values)
                good += 1
            except sqlite3.Error as exc:
                skipped.append({"line": line_no, "error": str(exc)})
        if not dry_run:
            target.commit()
        return {"inserted": good, "skipped": skipped, "table": table, "columns": insert_cols, "dry_run": dry_run}
    except (sqlite3.Error, ValueError, TypeError, KeyError, OSError):
        if not dry_run and target.in_transaction:
            target.rollback()
        raise
    finally:
        if own:
            target.close()


class AutomationEventBus:
    """Tiny in-process event bus for post-save automation."""
    def __init__(self) -> None:
        self._handlers: dict[str, list[Callable[[dict[str, Any]], None]]] = {}

    def subscribe(self, event: str, handler: Callable[[dict[str, Any]], None]) -> None:
        self._handlers.setdefault(event, []).append(handler)

    def publish(self, event: str, payload: dict[str, Any]) -> None:
        for handler in tuple(self._handlers.get(event, ())):
            try:
                handler(payload)
            except (ValueError, TypeError, KeyError, OSError, sqlite3.Error):
                # Secondary automation must never break the primary action.
                logger.exception("Automation event handler failed", extra={"event": event})


_EVENTS = AutomationEventBus()


def _audit_saved(payload: dict[str, Any]) -> None:
    try:
        from services.audit_service import audit
        audit("Auto-save", str(payload.get("entity", "Entity")), payload.get("entity_id"), "Created or updated through BuildOffice automation-aware workflow.")
    except (ImportError, ValueError, TypeError, KeyError, OSError, sqlite3.Error):
        logger.exception("Automation audit hook failed")


def _notify_status_change(payload: dict[str, Any]) -> None:
    try:
        new_status = str(payload.get("new_status") or "")
        if new_status.lower() in {"completed", "approved", "paid"}:
            conn = get_connection()
            try:
                conn.execute(
                    "INSERT INTO notifications(title,message,severity,entity,entity_id) VALUES(?,?,?,?,?)",
                    (f"{payload.get('entity', 'Record')} status updated", f"Status changed to {new_status}.", "Info", str(payload.get("entity", "Record")), str(payload.get("entity_id", ""))),
                )
                conn.commit()
            finally:
                conn.close()
    except (ValueError, TypeError, KeyError, OSError, sqlite3.Error):
        logger.exception("Automation status notification failed")


_EVENTS.subscribe("Project.saved", _audit_saved)
_EVENTS.subscribe("Expense.saved", _audit_saved)
_EVENTS.subscribe("Project.status_changed", _notify_status_change)
_EVENTS.subscribe("Situation.status_changed", _notify_status_change)
_EVENTS.subscribe("Supplementary.status_changed", _notify_status_change)


def events() -> AutomationEventBus:
    return _EVENTS


def emit_saved(entity: str, entity_id: str, **payload: Any) -> None:
    _EVENTS.publish(f"{entity}.saved", {"entity": entity, "entity_id": entity_id, **payload})


def emit_status_changed(entity: str, entity_id: str, old_status: str | None, new_status: str, **payload: Any) -> None:
    _EVENTS.publish(f"{entity}.status_changed", {"entity": entity, "entity_id": entity_id, "old_status": old_status, "new_status": new_status, **payload})


def automation_snapshot() -> dict[str, Any]:
    conn = get_connection()
    try:
        counts = {}
        for table in ("projects", "expenses", "expense_payments", "staff", "attendance", "vehicles", "documents", "debt"):
            try:
                counts[table] = int(conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0])
            except sqlite3.Error:
                counts[table] = 0
        return {"generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "counts": counts}
    finally:
        conn.close()


def generate_one_click_report(kind: str, project_id: str | None = None, fmt: str = "CSV") -> list[Path]:
    """Use existing report generators as the single-click rendering backend."""
    from services.document_generation_service import generate_text_or_csv
    return generate_text_or_csv(kind, project_id=project_id, fmt=fmt)


def export_automation_manifest(path: str | Path) -> Path:
    p = Path(path)
    p.write_text(json.dumps(automation_snapshot(), ensure_ascii=False, indent=2), encoding="utf-8")
    return p
