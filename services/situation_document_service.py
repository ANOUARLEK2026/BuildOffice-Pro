# ============================================================
# شرح الملف: services/situation_document_service.py
# هذا الملف جزء من منطق الأعمال والخدمات والتحليلات.
# التعليقات العربية هنا للتوضيح والصيانة فقط، ولا تغيّر سلوك البرنامج.
# تم الحفاظ على أسماء الكائنات وواجهات الاستدعاء ومنطق قاعدة البيانات كما هي.
# ============================================================
"""Generate Word and Excel document packages for a project situation."""
from __future__ import annotations

import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from core.config import situation_generated_dir
from core.id_service import next_child_id
from repositories.db import get_connection
from services.settings_service import get_setting

from .situation_service import get_situation

SITUATION_OUTPUT_DIR = situation_generated_dir()


# شرح الدالة _require_dependencies: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _require_dependencies():
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT, WD_SECTION
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, Side
    except ImportError as exc:
        missing = []
        try:
            from services.dependency_service import missing_document_dependencies
            missing = missing_document_dependencies()
        except (ImportError, AttributeError, TypeError):
            pass
        package_text = ", ".join(missing) if missing else "python-docx and openpyxl"
        raise RuntimeError(
            "Word/Excel generation cannot start because these packages are missing: "
            f"{package_text}.\n\n"
            "Run 'Install BuildOffice Pro Dependencies.bat' in the application folder, "
            "then restart BuildOffice Pro."
        ) from exc
    return {
        "Document": Document,
        "WD_SECTION": WD_SECTION,
        "WD_ORIENT": WD_ORIENT,
        "Inches": Inches,
        "Pt": Pt,
        "WD_ALIGN_PARAGRAPH": WD_ALIGN_PARAGRAPH,
        "OxmlElement": OxmlElement,
        "qn": qn,
        "Workbook": Workbook,
        "Font": Font,
        "Alignment": Alignment,
        "Border": Border,
        "Side": Side,
    }


# شرح الدالة _money: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _money(value: Any) -> str:
    return f"{float(value or 0):,.2f}"


# شرح الدالة _qty: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _qty(value: Any) -> str:
    return f"{float(value or 0):,.3f}"


# شرح الدالة _language: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _language() -> str:
    value = str(get_setting("language", "en")).strip().lower()
    return "ar" if value in ("ar", "arabic") else "en"


# شرح الدالة _tr: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _tr(language: str, en: str, ar: str) -> str:
    return ar if language == "ar" else en


# شرح الدالة _set_doc_rtl: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _set_doc_rtl(paragraph, deps):
    if _language() != "ar":
        return
    paragraph.alignment = deps["WD_ALIGN_PARAGRAPH"].RIGHT


# شرح الدالة _set_table_rtl: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _set_table_rtl(table, deps):
    if _language() != "ar":
        return
    tbl_pr = table._tbl.tblPr
    bidi = deps["OxmlElement"]("w:bidiVisual")
    bidi.set(deps["qn"]("w:val"), "1")
    tbl_pr.append(bidi)


# شرح الدالة _configure_doc: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _configure_doc(document, deps):
    section = document.sections[0]
    section.top_margin = deps["Inches"](0.55)
    section.bottom_margin = deps["Inches"](0.55)
    section.left_margin = deps["Inches"](0.55)
    section.right_margin = deps["Inches"](0.55)
    return document


# شرح الدالة _add_doc_title: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _add_doc_title(document, title: str, subtitle: str, deps):
    p = document.add_paragraph()
    p.alignment = deps["WD_ALIGN_PARAGRAPH"].CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = deps["Pt"](16)
    p2 = document.add_paragraph(subtitle)
    p2.alignment = deps["WD_ALIGN_PARAGRAPH"].CENTER


# شرح الدالة _add_project_meta: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _add_project_meta(document, situation: dict[str, Any], deps):
    lang = _language()
    s = situation["situation"]
    lines = [
        (_tr(lang, "Project", "المشروع"), s["project_name"]),
        (_tr(lang, "Contract code", "رقم العقد"), s.get("contract_code") or ""),
        (_tr(lang, "Authority", "صاحب المشروع"), s.get("contracting_authority") or ""),
        (_tr(lang, "Situation", "الوضعية"), f"N°{s['situation_number']}"),
        (_tr(lang, "Date", "التاريخ"), s["situation_date"]),
        (_tr(lang, "Status", "الحالة"), s["status"]),
    ]
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    _set_table_rtl(table, deps)
    for key, value in lines:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)


# شرح الدالة _add_totals_table: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _add_totals_table(document, situation: dict[str, Any], deps):
    lang = _language()
    labels = [
        (_tr(lang, "Contract / BOQ value", "قيمة العقد / جدول الكميات"), situation["boq_value"]),
        (_tr(lang, "Supplementary works", "الأشغال الإضافية"), sum(float(x.get("current_amount") or 0) for x in situation.get("supplementary_items", []))),
        (_tr(lang, "Previous cumulative", "المنجز السابق التراكمي"), situation["previous_value"]),
        (_tr(lang, "Current situation HT", "قيمة الوضعية الحالية HT"), situation["period_value"]),
        (_tr(lang, "Current TVA", "TVA الوضعية الحالية"), situation["period_tva"]),
        (_tr(lang, "Current situation TTC", "قيمة الوضعية الحالية TTC"), situation["period_ttc"]),
        (_tr(lang, "Cumulative executed HT", "المنجز التراكمي HT"), situation["cumulative_value"]),
        (_tr(lang, "Cumulative TVA", "TVA التراكمي"), situation["cumulative_tva"]),
        (_tr(lang, "Cumulative TTC", "المنجز التراكمي TTC"), situation["cumulative_ttc"]),
        (_tr(lang, "Remaining HT", "المتبقي HT"), situation["remaining_value"]),
        (_tr(lang, "Progress", "نسبة التقدم"), f"{situation['project_progress_percent']:.2f}%"),
    ]
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    _set_table_rtl(table, deps)
    for label, value in labels:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value) if isinstance(value, str) else _money(value)


# شرح الدالة _build_summary_word: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _build_summary_word(situation: dict[str, Any], path: Path, deps):
    lang = _language()
    document = _configure_doc(deps["Document"](), deps)
    _add_doc_title(
        document,
        _tr(lang, "Construction Work Situation", "وضعية أشغال البناء والأشغال العمومية"),
        _tr(lang, "Generated by BuildOffice Pro", "تم إنشاؤها بواسطة BuildOffice Pro"),
        deps,
    )
    _add_project_meta(document, situation, deps)
    document.add_paragraph()
    _add_totals_table(document, situation, deps)
    if situation["situation"].get("notes"):
        p = document.add_paragraph(
            f"{_tr(lang, 'Notes', 'ملاحظات')}: {situation['situation']['notes']}"
        )
        _set_doc_rtl(p, deps)
    document.save(path)


# شرح الدالة _build_detail_word: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _build_detail_word(situation: dict[str, Any], path: Path, deps):
    lang = _language()
    document = _configure_doc(deps["Document"](), deps)
    section = document.sections[0]
    section.orientation = deps["WD_ORIENT"].LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    _add_doc_title(
        document,
        _tr(lang, "Situation Detailed Statement", "الكشف التفصيلي للوضعية"),
        f"N°{situation['situation']['situation_number']} • {situation['situation']['situation_date']}",
        deps,
    )
    _add_project_meta(document, situation, deps)
    document.add_paragraph()

    headers = [
        _tr(lang, "No.", "الرقم"),
        _tr(lang, "Description", "البيان"),
        _tr(lang, "Unit", "الوحدة"),
        _tr(lang, "Contract Qty", "كمية العقد"),
        _tr(lang, "Previous", "السابق"),
        _tr(lang, "Current", "الحالي"),
        _tr(lang, "Cumulative", "التراكمي"),
        _tr(lang, "Remaining", "المتبقي"),
        _tr(lang, "Unit Price", "سعر الوحدة"),
        _tr(lang, "Current Amount", "مبلغ الوضعية"),
    ]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    _set_table_rtl(table, deps)
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for item in situation["items"]:
        row = table.add_row().cells
        values = [
            item["item_number"] or "-",
            item["description"],
            item["unit"] or "",
            _qty(item["contract_quantity"]),
            _qty(item["previous_quantity"]),
            _qty(item["current_quantity"]),
            _qty(item["cumulative_quantity"]),
            _qty(item["remaining_quantity"]),
            _money(item["unit_price"]),
            _money(item["current_amount"]),
        ]
        for i, value in enumerate(values):
            row[i].text = str(value)

    if situation.get("supplementary_items"):
        document.add_paragraph(_tr(lang, "Supplementary / Additional Works", "الأشغال الإضافية"))
        st = document.add_table(rows=1, cols=7); st.style = "Table Grid"; _set_table_rtl(st, deps)
        for i,h in enumerate([_tr(lang,"No.","الرقم"),_tr(lang,"Description","البيان"),_tr(lang,"Unit","الوحدة"),_tr(lang,"Mode","الطريقة"),_tr(lang,"Quantity","الكمية"),_tr(lang,"Unit Price","سعر الوحدة"),_tr(lang,"Amount","المبلغ")]): st.rows[0].cells[i].text=h
        for x in situation["supplementary_items"]:
            c=st.add_row().cells; vals=[f"SW-{int(x['work_number']):03d}",x["description"],x.get("unit") or "",x.get("work_mode") or "",_qty(x.get("current_quantity")),_money(x.get("unit_price")),_money(x.get("current_amount"))]
            for i,v in enumerate(vals): c[i].text=str(v)
    document.add_paragraph()
    _add_totals_table(document, situation, deps)
    document.save(path)


# شرح الدالة _style_excel: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _style_excel(ws, deps, rtl: bool):
    thin = deps["Side"](style="thin")
    border = deps["Border"](left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = deps["Alignment"](vertical="center", wrap_text=True)
    if rtl:
        ws.sheet_view.rightToLeft = True


# شرح الدالة _build_summary_excel: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _build_summary_excel(situation: dict[str, Any], path: Path, deps):
    lang = _language()
    wb = deps["Workbook"]()
    ws = wb.active
    ws.title = _tr(lang, "Summary", "الملخص")
    ws.append([_tr(lang, "Construction Work Situation", "وضعية الأشغال")])
    ws.append([_tr(lang, "Project", "المشروع"), situation["situation"]["project_name"]])
    ws.append([_tr(lang, "Situation", "الوضعية"), situation["situation"]["situation_number"]])
    ws.append([_tr(lang, "Date", "التاريخ"), situation["situation"]["situation_date"]])
    ws.append([_tr(lang, "Status", "الحالة"), situation["situation"]["status"]])
    ws.append([])
    rows = [
        (_tr(lang, "Contract / BOQ value", "قيمة العقد / جدول الكميات"), situation["boq_value"]),
        (_tr(lang, "Previous cumulative", "المنجز السابق التراكمي"), situation["previous_value"]),
        (_tr(lang, "Current situation HT", "قيمة الوضعية الحالية HT"), situation["period_value"]),
        (_tr(lang, "Current TVA", "TVA الوضعية الحالية"), situation["period_tva"]),
        (_tr(lang, "Current situation TTC", "قيمة الوضعية الحالية TTC"), situation["period_ttc"]),
        (_tr(lang, "Cumulative executed HT", "المنجز التراكمي HT"), situation["cumulative_value"]),
        (_tr(lang, "Cumulative TVA", "TVA التراكمي"), situation["cumulative_tva"]),
        (_tr(lang, "Cumulative TTC", "المنجز التراكمي TTC"), situation["cumulative_ttc"]),
        (_tr(lang, "Remaining HT", "المتبقي HT"), situation["remaining_value"]),
        (_tr(lang, "Progress %", "نسبة التقدم"), situation["project_progress_percent"] / 100),
    ]
    for row in rows:
        ws.append(list(row))
    ws["A1"].font = deps["Font"](bold=True, size=16)
    for row_index in range(7, 16):
        ws[f"B{row_index}"].number_format = '#,##0.00'
    ws["B16"].number_format = '0.00%'
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 38
    _style_excel(ws, deps, lang == "ar")
    wb.save(path)


# شرح الدالة _build_detail_excel: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _build_detail_excel(situation: dict[str, Any], path: Path, deps):
    lang = _language()
    wb = deps["Workbook"]()
    ws = wb.active
    ws.title = _tr(lang, "Situation detail", "تفصيل الوضعية")
    headers = [
        _tr(lang, "No.", "الرقم"),
        _tr(lang, "Description", "البيان"),
        _tr(lang, "Unit", "الوحدة"),
        _tr(lang, "Contract Qty", "كمية العقد"),
        _tr(lang, "Previous", "السابق"),
        _tr(lang, "Current", "الحالي"),
        _tr(lang, "Cumulative", "التراكمي"),
        _tr(lang, "Remaining", "المتبقي"),
        _tr(lang, "Unit Price", "سعر الوحدة"),
        _tr(lang, "Current Amount", "مبلغ الوضعية"),
    ]
    ws.append(headers)
    for item in situation["items"]:
        ws.append([
            item["item_number"] or "-",
            item["description"],
            item["unit"] or "",
            item["contract_quantity"],
            item["previous_quantity"],
            item["current_quantity"],
            item["cumulative_quantity"],
            item["remaining_quantity"],
            item["unit_price"],
            item["current_amount"],
        ])
    widths = [12, 48, 14, 15, 15, 15, 15, 15, 18, 20]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=8):
        for cell in row:
            cell.number_format = '#,##0.000'
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=10):
        for cell in row:
            cell.number_format = '#,##0.00'
    ws.freeze_panes = "A2"
    _style_excel(ws, deps, lang == "ar")
    wb.save(path)


# شرح الدالة _register_generated_documents: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def _register_generated_documents(project_id: str, situation_id: str, files: list[Path]) -> None:
    """Register generated files in both the dedicated table and existing project documents."""
    conn = get_connection()
    try:
        conn.execute("BEGIN IMMEDIATE")
        conn.execute("DELETE FROM situation_documents WHERE situation_id=?", (situation_id,))
        conn.execute(
            'DELETE FROM "projects documents" WHERE related_entity=?',
            (situation_id,),
        )
        for path in files:
            document_type = path.suffix.lower().lstrip(".")
            doc_id = next_child_id(conn, "sdoc", situation_id, digits=3)
            conn.execute(
                """
                INSERT INTO situation_documents(
                    document_id, situation_id, project_id, document_name,
                    document_path, document_type, generated_at
                ) VALUES(?,?,?,?,?,?,CURRENT_TIMESTAMP)
                """,
                (doc_id, situation_id, project_id, path.name, str(path), document_type),
            )
            project_doc_id = next_child_id(conn, "pdoc", project_id, digits=4)
            conn.execute(
                """
                INSERT INTO "projects documents"(
                    item_id, project_id, document_name, document_path,
                    upload_date, description, status, category, nature, related_entity
                ) VALUES(?,?,?,?,CURRENT_TIMESTAMP,?,'Active',?,?,?)
                """,
                (
                    project_doc_id,
                    project_id,
                    path.name,
                    str(path),
                    f"Generated from situation {situation_id}",
                    "Situation Package",
                    "Situation",
                    situation_id,
                ),
            )
        conn.commit()
    except (sqlite3.Error, ValueError, TypeError, OSError):
        conn.rollback()
        raise
    finally:
        conn.close()


# شرح الدالة generate_situation_package: تنفّذ خطوة منطقية مرتبطة بوظيفة هذا الملف.
def generate_situation_package(situation_id: str, output_root: str | Path | None = None) -> dict[str, Any]:
    """Generate the standard Word/Excel package and return its paths."""
    situation = get_situation(situation_id)
    if not situation:
        raise ValueError("Situation not found.")

    deps = _require_dependencies()
    root = Path(output_root) if output_root else SITUATION_OUTPUT_DIR
    project_id = str(situation["situation"]["project_id"])
    number = int(situation["situation"]["situation_number"])
    folder = root / project_id / f"Situation-{number:03d}"
    folder.mkdir(parents=True, exist_ok=True)

    files = [
        folder / f"01_Situation-{number:03d}_Summary.docx",
        folder / f"02_Situation-{number:03d}_Detail.docx",
        folder / f"03_Situation-{number:03d}_Summary.xlsx",
        folder / f"04_Situation-{number:03d}_Detail.xlsx",
    ]
    _build_summary_word(situation, files[0], deps)
    _build_detail_word(situation, files[1], deps)
    _build_summary_excel(situation, files[2], deps)
    _build_detail_excel(situation, files[3], deps)
    _register_generated_documents(project_id, situation_id, files)

    return {
        "situation_id": situation_id,
        "folder": folder,
        "files": files,
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }


__all__ = ["SITUATION_OUTPUT_DIR", "generate_situation_package"]
